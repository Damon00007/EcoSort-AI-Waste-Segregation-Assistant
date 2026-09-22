from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 1. Initialize presentation (16:9 widescreen)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# 2. Color Palette Definitions
DARK_BG = RGBColor(13, 40, 24)       # Deep Forest Green
CARD_BG = RGBColor(22, 56, 38)       # Translucent Dark Card
EMERALD = RGBColor(0, 230, 118)      # Accent Green
AMBER = RGBColor(255, 183, 3)        # Highlight Gold
WHITE = RGBColor(255, 255, 255)
LIGHT_GRAY = RGBColor(220, 220, 220)

def set_background(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_BG
    bg.line.fill.background()
    return bg

def add_header(slide, category, title):
    # Category / Pill header
    tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
    p_cat = tb_cat.text_frame.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = EMERALD

    # Main Slide Title
    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
    p_title = tb_title.text_frame.paragraphs[0]
    p_title.text = title
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE

def add_card(slide, left, top, width, height, border_color=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
    else:
        card.line.fill.background()
    return card

# ==========================================
# SLIDE 1: TITLE SLIDE
# ==========================================
slide1 = prs.slides.add_slide(blank_layout)
set_background(slide1)

# Badge
badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(6.5), Inches(0.5))
badge.fill.solid()
badge.fill.fore_color.rgb = CARD_BG
badge.line.color.rgb = EMERALD
p_b = badge.text_frame.paragraphs[0]
p_b.text = "1M1B × IBM SkillsBuild × AICTE Internship Project"
p_b.font.size = Pt(12)
p_b.font.bold = True
p_b.font.color.rgb = EMERALD
p_b.alignment = PP_ALIGN.CENTER

# Main Title & Subtitle
tb = slide1.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.5), Inches(2.2))
tf = tb.text_frame
p1 = tf.paragraphs[0]
p1.text = "EcoSort AI"
p1.font.size = Pt(50)
p1.font.bold = True
p1.font.color.rgb = WHITE

p2 = tf.add_paragraph()
p2.text = "Intelligent Source-Level Waste Segregation & Disposal Assistant"
p2.font.size = Pt(22)
p2.font.color.rgb = EMERALD

# Student Details Card
c_user = add_card(slide1, Inches(0.8), Inches(4.7), Inches(6.5), Inches(1.8), EMERALD)
tf_u = c_user.text_frame
tf_u.margin_left = Inches(0.3)
tf_u.margin_top = Inches(0.25)
p_u1 = tf_u.paragraphs[0]
p_u1.text = "Submitted by: Pawan Tiwari"
p_u1.font.size = Pt(18)
p_u1.font.bold = True
p_u1.font.color.rgb = WHITE

p_u2 = tf_u.add_paragraph()
p_u2.text = "B.Tech Computer Science & Engineering\nAmbalika Institute of Management & Technology"
p_u2.font.size = Pt(13)
p_u2.font.color.rgb = LIGHT_GRAY

# ==========================================
# SLIDE 2: SDG ALIGNMENT & PROBLEM STATEMENT
# ==========================================
slide2 = prs.slides.add_slide(blank_layout)
set_background(slide2)
add_header(slide2, "Context & Goals", "SDG Alignment & Ground Reality Problem")

# Left: Problem
c_prob = add_card(slide2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(3.8))
tf_p = c_prob.text_frame
tf_p.margin_left = Inches(0.3)
tf_p.margin_top = Inches(0.3)
p = tf_p.paragraphs[0]
p.text = "The Ground Reality"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = AMBER

p = tf_p.add_paragraph()
p.text = "\n• Source Mixing: Wet food waste, clean recyclables, and toxic cells are commingled into single trash containers.\n• Landfill Crisis: Contamination prevents up to 70% of recyclable materials from processing, leading to heavy methane emissions.\n• Health Hazards: Manual handling of mixed waste exposes sanitation workers to direct chemical and physical risks."
p.font.size = Pt(13)
p.font.color.rgb = LIGHT_GRAY

# Right: SDG Mapping
c_sdg = add_card(slide2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(3.8))
tf_s = c_sdg.text_frame
tf_s.margin_left = Inches(0.3)
tf_s.margin_top = Inches(0.3)
p = tf_s.paragraphs[0]
p.text = "UN SDG Framework"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = EMERALD

p = tf_s.add_paragraph()
p.text = "\n• Primary Goal: SDG 12 (Responsible Consumption & Production)\n  - Sub-Target 12.5: Substantially reduce waste generation through prevention, reduction, recycling, and reuse.\n\n• Secondary Goal: SDG 11 (Sustainable Cities & Communities)\n  - Target 11.6: Reduce adverse per capita environmental impact of municipal waste."
p.font.size = Pt(13)
p.font.color.rgb = LIGHT_GRAY

# HMW Statement Bar
c_hmw = add_card(slide2, Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.3), EMERALD)
tf_h = c_hmw.text_frame
tf_h.margin_left = Inches(0.3)
tf_h.margin_top = Inches(0.2)
p = tf_h.paragraphs[0]
p.text = "How Might We (HMW) Design Statement:"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = EMERALD
p = tf_h.add_paragraph()
p.text = '"How might we use AI to guide everyday users in segregating waste at the exact point of disposal so that material recovery and recycling become effortless and accurate?"'
p.font.size = Pt(13)
p.font.color.rgb = WHITE

# ==========================================
# SLIDE 3: TARGET USERS & THE NEED FOR AI
# ==========================================
slide3 = prs.slides.add_slide(blank_layout)
set_background(slide3)
add_header(slide3, "Stakeholders & Justification", "Target Users & Why AI is Essential")

users = [
    ("Target Users", "• College students, faculty & hostelites\n• Campus cafeteria & canteen teams\n• Residential housing societies\n• Municipal & campus sanitation workers", EMERALD),
    ("Limitations of Static Bins", "• Static color stickers fail on composite items (e.g., plastic-lined paper coffee cups).\n• Zero user awareness on prep steps like rinsing oily residues or taping battery ends.", AMBER),
    ("The AI Advantage", "• Real-Time Resolution: Instantly classifies complex items into exact bin colors.\n• Habit Formation: Delivers prep steps (rinse/crush) + environmental feedback metrics.", EMERALD)
]

for idx, (title, content, color) in enumerate(users):
    card = add_card(slide3, Inches(0.8 + idx * 4.0), Inches(1.8), Inches(3.7), Inches(4.8))
    tf = card.text_frame
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = color
    p2 = tf.add_paragraph()
    p2.text = f"\n{content}"
    p2.font.size = Pt(13)
    p2.font.color.rgb = LIGHT_GRAY

# ==========================================
# SLIDE 4: SYSTEM ARCHITECTURE & WORKFLOW
# ==========================================
slide4 = prs.slides.add_slide(blank_layout)
set_background(slide4)
add_header(slide4, "System Pipeline", "EcoSort AI Architecture & Data Flow")

steps = [
    ("Step 1: Input Layer", "User provides input via text query or image capture (e.g., 'plastic bottle')."),
    ("Step 2: AI Engine", "IBM Granite / Multimodal LLM extracts material profile (PET, Cardboard, Organic)."),
    ("Step 3: Logic Engine", "Evaluates contamination level and cross-references municipal bin regulations."),
    ("Step 4: Action Output", "Displays bin color, mandatory pre-disposal steps, and sustainability savings.")
]

for idx, (title, content) in enumerate(steps):
    card = add_card(slide4, Inches(0.8 + idx * 3.0), Inches(2.2), Inches(2.7), Inches(3.8), EMERALD if idx == 1 else None)
    tf = card.text_frame
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = EMERALD if idx == 1 else WHITE
    p2 = tf.add_paragraph()
    p2.text = f"\n{content}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = LIGHT_GRAY

# ==========================================
# SLIDE 5: RESPONSIBLE AI CONSIDERATIONS
# ==========================================
slide5 = prs.slides.add_slide(blank_layout)
set_background(slide5)
add_header(slide5, "Ethics & Safety", "Responsible AI Considerations")

ethics = [
    ("Fairness & Localization", "Accurately identifies localized Indian waste items (kulhads, coconut husks, snack sachets) without data bias.", Inches(0.8), Inches(1.8)),
    ("Explainability", "Clearly explains the rationale behind every sorting decision to instill long-term user segregation habits.", Inches(6.8), Inches(1.8)),
    ("Privacy-First Design", "Operates exclusively on object imagery; completely avoids processing human faces or personal identifiers.", Inches(0.8), Inches(4.3)),
    ("Safety & Hazards", "Triggers high-priority visual caution alerts whenever sharp glass, e-waste, or batteries are detected.", Inches(6.8), Inches(4.3))
]

for title, desc, left, top in ethics:
    c = add_card(slide5, left, top, Inches(5.7), Inches(2.2))
    tf = c.text_frame
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p2 = tf.add_paragraph()
    p2.text = f"\n{desc}"
    p2.font.size = Pt(13)
    p2.font.color.rgb = LIGHT_GRAY

# ==========================================
# SLIDE 6: PROTOTYPE DEMONSTRATION
# ==========================================
slide6 = prs.slides.add_slide(blank_layout)
set_background(slide6)
add_header(slide6, "Implementation", "Prototype Demonstration (Streamlit Web App)")

# Left: Placeholder card for Screenshot
c_img = add_card(slide6, Inches(0.8), Inches(1.8), Inches(6.8), Inches(4.8), EMERALD)
tf_i = c_img.text_frame
tf_i.margin_left = Inches(0.3)
tf_i.margin_top = Inches(0.3)
p = tf_i.paragraphs[0]
p.text = "[ Streamlit Demo Screenshot Container ]"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = EMERALD
p2 = tf_i.add_paragraph()
p2.text = "\n• Working web app deployed locally with interactive classification.\n• Crop & paste your Streamlit screenshot over this card."
p2.font.size = Pt(13)
p2.font.color.rgb = LIGHT_GRAY

# Right: Validated Test Cases
c_cases = add_card(slide6, Inches(7.9), Inches(1.8), Inches(4.6), Inches(4.8))
tf_c = c_cases.text_frame
tf_c.margin_left = Inches(0.3)
tf_c.margin_top = Inches(0.3)
p = tf_c.paragraphs[0]
p.text = "Validated Test Cases"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = WHITE

cases_text = [
    ("🍌 Banana Peel", "🟢 Green Bin (Wet/Compost)\nAction: Discard directly without polybag."),
    ("🧴 Plastic Bottle", "🔵 Blue Bin (Dry Recyclable)\nAction: Empty liquid, rinse & crush."),
    ("🔋 Lithium Cell", "🔴 Red/Black Bin (Hazardous E-Waste)\nAction: Tape terminals & deposit at e-kiosk.")
]
for title, desc in cases_text:
    p_t = tf_c.add_paragraph()
    p_t.text = f"\n{title}"
    p_t.font.size = Pt(14)
    p_t.font.bold = True
    p_t.font.color.rgb = AMBER
    p_d = tf_c.add_paragraph()
    p_d.text = desc
    p_d.font.size = Pt(12)
    p_d.font.color.rgb = LIGHT_GRAY

# ==========================================
# SLIDE 7: EXPECTED IMPACT & FUTURE SCOPE
# ==========================================
slide7 = prs.slides.add_slide(blank_layout)
set_background(slide7)
add_header(slide7, "Outcomes & Roadmap", "Measurable Impact & Future Scope")

impacts = [
    ("Environmental Impact", "• Landfill Diversion: Significantly lowers landfill load and stops methane formation.\n• Circular Value: Improves yield and quality of recycled plastics and paper."),
    ("Social & Occupational Dignity", "• Worker Safety: Protects sanitation staff from hazardous cuts and chemical exposure.\n• Community Pride: Enhances civic responsibility across campus corridors."),
    ("Campus Benchmarking & Future", "• Green Audit: Helps institutions secure higher ratings in environmental certifications.\n• Future IoT Scope: Integration with camera-enabled smart bin lids for automatic mechanical sorting.")
]

for idx, (title, content) in enumerate(impacts):
    c = add_card(slide7, Inches(0.8 + idx * 4.0), Inches(1.8), Inches(3.7), Inches(4.8))
    tf = c.text_frame
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p2 = tf.add_paragraph()
    p2.text = f"\n{content}"
    p2.font.size = Pt(13)
    p2.font.color.rgb = LIGHT_GRAY

# 3. Save the generated presentation
prs.save("EcoSort_AI_Professional.pptx")
print("Successfully generated: EcoSort_AI_Professional.pptx")