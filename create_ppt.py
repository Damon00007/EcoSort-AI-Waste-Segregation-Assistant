from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()

# Slide 1: Title
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "EcoSort AI – Smart Waste Segregation Assistant"
subtitle.text = "1M1B AI for Sustainability Virtual Internship\nIn Collaboration with IBM SkillsBuild & AICTE\n\nSubmitted by: Pawan Tiwari\nAmbalika Institute of Management & Technology"

# Function to add content slides
def add_custom_slide(prs, title_text, bullet_points):
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = title_text
    tf = body_shape.text_frame
    tf.word_wrap = True

    for i, pt in enumerate(bullet_points):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = pt
        p.level = 0
        p.font.size = Pt(18)

# Slide 2
add_custom_slide(prs, "SDG Alignment & Problem Statement", [
    "Primary SDG: SDG 12 (Responsible Consumption & Production - Target 12.5).",
    "Secondary SDG: SDG 11 (Sustainable Cities and Communities).",
    "Problem: Commingling of wet, dry recyclable, and hazardous waste at source causes recyclable materials to rot in landfills.",
    "HMW Question: How might we use AI to guide users in segregating waste at the exact point of disposal to maximize recycling?"
])

# Slide 3
add_custom_slide(prs, "Target Users & The Need for AI", [
    "Target Users: College students, campus cafeteria staff, residential societies, and municipal waste workers.",
    "Speed & Real-time Precision: Instantly resolves ambiguity for composite items (e.g., plastic-lined paper tea cups).",
    "Multimodal Scalability: Supports text and vision-based item verification without requiring complex manuals.",
    "Behavioral Impact: Encourages pre-disposal preparation (rinsing, crushing, separating components)."
])

# Slide 4
add_custom_slide(prs, "AI Solution & System Architecture", [
    "Input Layer: User enters waste item name or uploads photo.",
    "Processing Engine: IBM Granite / Multimodal LLM extracts material properties (PET, Paper, Organic, Hazardous).",
    "Rules Engine: Maps material against municipal color-coded bin standards (Green, Blue, Red/Black).",
    "Output Generation: Delivers specific bin color, cleaning steps, and sustainability impact."
])

# Slide 5
add_custom_slide(prs, "Responsible AI Considerations", [
    "Fairness: Calibrated for Indian household and campus waste (e.g., kulhad, coconut husks, snack pouches).",
    "Explainability: Explains the 'Why' behind bin allocation to build user habits.",
    "Privacy: Operates strictly on object data; captures zero user faces or personal identifiers.",
    "Safety: Flags sharp glass, chemical containers, and batteries with explicit handling warnings."
])

# Slide 6
add_custom_slide(prs, "Prototype Demonstration", [
    "Demo Interface: Interactive Streamlit Web Application (EcoSort AI).",
    "Tested Case 1 (Plastic Bottle): Blue Bin (Dry Recyclable) -> Action: Rinse & Crush.",
    "Tested Case 2 (Banana Peel): Green Bin (Wet/Compost) -> Action: Dispose without plastic bag.",
    "Tested Case 3 (Lithium Battery): Red/Black Bin (Hazardous E-Waste) -> Action: Tape terminals.",
    "[Paste your Streamlit app screenshots here]"
])

# Slide 7
add_custom_slide(prs, "Expected Impact & Future Scope", [
    "Environmental: Significant reduction in landfill burden and methane generation from mixed piles.",
    "Social Dignity: Shields sanitation workers from manual contact with broken glass and toxic cells.",
    "Campus Metric: Enables colleges to meet Green Audit and Swachh Campus benchmarks.",
    "Future Scope: IoT integration with camera-equipped smart dustbin lids for automatic mechanical sorting."
])

# Save File
prs.save("EcoSort_AI_1M1B_Project.pptx")
print("Presentation generated successfully: EcoSort_AI_1M1B_Project.pptx")